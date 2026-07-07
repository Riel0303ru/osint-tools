# utils/result_formatter.py
from typing import List, Dict, Optional, Union
from rich.console import Console
from rich.table import Table
from openpyxl import Workbook
import os
import csv
import re
from core.base.scan_result import ScanResult

console = Console()


class ResultFormatter:

    def __init__(self, base_output_dir: str = "Osint/results"):
        self.base_output_dir = base_output_dir

    @staticmethod
    def _sanitize_filename(name: str) -> str:
        """Ganti karakter ilegal Windows/Unix dengan underscore."""
        return re.sub(r'[<>:"/\\|?*]', '_', name)
    
    def _get_user_dir(self, username: str, category: Optional[str] = None) -> str:
        safe_username = self._sanitize_filename(username)
        if category:
            target_dir = os.path.join(self.base_output_dir, category, safe_username)
        else:
            target_dir = os.path.join(self.base_output_dir, safe_username)
        os.makedirs(target_dir, exist_ok=True)
        return target_dir

    def _normalize_results(self, results: List[Union[ScanResult, Dict]]) -> List[Dict]:
        normalized = []
        for r in results:
            if isinstance(r, ScanResult):
                normalized.append(r.to_dict())
            elif isinstance(r, dict):
                normalized.append(r)
            else:
                normalized.append({"platform": "unknown", "error": "invalid result"})
        return normalized

    def group_by_platform(self, results: List[Union[ScanResult, Dict]]) -> Dict[str, List[Dict]]:
        grouped = {}
        normalized = self._normalize_results(results)
        for r in normalized:
            platform = r.get("platform", "Unknown")
            grouped.setdefault(platform, []).append(r)
        return grouped
    def export_csv_per_platform(
        self,
        results: List[Union[ScanResult, Dict]],
        username: str,
        category: Optional[str] = None
    ):
        target_dir = self._get_user_dir(username, category)
        grouped = self.group_by_platform(results)
        headers = ["Platform", "Username", "Status", "Status Code", "URL", "Confidence", "Intelligence Score"]
        for platform, items in grouped.items():
            file_path = os.path.join(target_dir, f"{platform}.csv")
            with open(file_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for r in items:
                    writer.writerow([
                        r.get("platform", ""),
                        r.get("username", ""),
                        r.get("status", ""),
                        r.get("status_code", ""),
                        r.get("url", ""),
                        r.get("confidence", 1.0),
                        r.get("intelligence_score", 0.0)
                    ])
    def export_excel_summary(
        self,
        results: List[Union[ScanResult, Dict]],
        username: str,
        category: Optional[str] = None
    ):
        target_dir = self._get_user_dir(username, category)
        file_path = os.path.join(target_dir, "summary.xlsx")
        wb = Workbook()
        wb.remove(wb.active)

        normalized = self._normalize_results(results)
        grouped = self.group_by_platform(results)

        headers = ["Platform", "Username", "Status", "Status Code", "URL", "Confidence", "Intelligence Score"]

        # Summary sheet
        ws_summary = wb.create_sheet("Summary", 0)
        ws_summary.append(headers)
        for r in normalized:
            ws_summary.append([
                r.get("platform", ""),
                r.get("username", ""),
                r.get("status", ""),
                r.get("status_code", ""),
                r.get("url", ""),
                r.get("confidence", 1.0),
                r.get("intelligence_score", 0.0)
            ])

        # Per platform
        for platform, items in grouped.items():
            sheet_name = platform[:31]
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
            for r in items:
                ws.append([
                    r.get("platform", ""),
                    r.get("username", ""),
                    r.get("status", ""),
                    r.get("status_code", ""),
                    r.get("url", ""),
                    r.get("confidence", 1.0),
                    r.get("intelligence_score", 0.0)
                ])

        wb.save(file_path)

    def create_scan_table(self, results: List[Union[ScanResult, Dict]], title: str):
        table = Table(title=f"[bold cyan]{title}[/bold cyan]", border_style="cyan")
        table.add_column("Platform")
        table.add_column("Username")
        table.add_column("Status")
        table.add_column("Code")
        table.add_column("URL")
        table.add_column("Intel Score")
        table.add_column("Info")

        normalized = self._normalize_results(results)
        for r in normalized:
            extra = r.get("extra", {})
            info_indicator = "⭐" if extra else "-"
            table.add_row(
                r.get("platform", "-"),
                r.get("username", "-"),
                r.get("status", "-"),
                str(r.get("status_code", "-")),
                r.get("url", "-"),
                f"{r.get('intelligence_score', 0.0):.1f}",
                info_indicator
            )
        return table
    
    def display_full_report(
        self,
        results: List[Union[ScanResult, Dict]],
        title: str,
        username: str,
        category: Optional[str] = None,
        export: bool = True
    ):
        console.print(self.create_scan_table(results, title))

        if export and results:
            self.export_csv_per_platform(results, username, category)
            self.export_excel_summary(results, username, category)
            out_path = self._get_user_dir(username, category)
            console.print(f"\n[green]Results exported to {out_path}[/green]")