# core/document/extractors/xlsx_extractor.py
from __future__ import annotations

from typing import Dict, Any
import openpyxl


class XLSXExtractor:
    """Ekstraktor untuk file XLSX."""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        result = {
            "metadata": {},
            "sheets": [],
            "total_rows": 0,
            "text": "",
        }

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)

            # Metadata
            props = wb.properties
            result["metadata"] = {
                "creator": props.creator,
                "created": str(props.created) if props.created else None,
                "modified": str(props.modified) if props.modified else None,
                "last_modified_by": props.lastModifiedBy,
                "title": props.title,
                "subject": props.subject,
                "keywords": props.keywords,
            }

            # Sheets
            text_parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_count = ws.max_row or 0
                result["sheets"].append({"name": sheet_name, "rows": rows_count})
                result["total_rows"] += rows_count

                # Text dari semua sel
                for row in ws.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        text_parts.append(row_text)

            result["text"] = "\n".join(text_parts)
            wb.close()

        except Exception as e:
            result["error"] = str(e)

        return result