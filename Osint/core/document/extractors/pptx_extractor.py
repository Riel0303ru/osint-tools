# core/document/extractors/pptx_extractor.py
from __future__ import annotations

from typing import Dict, Any
from pptx import Presentation


class PPTXExtractor:
    """Ekstraktor untuk file PPTX."""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        result = {
            "metadata": {},
            "slides_count": 0,
            "text": "",
        }

        try:
            prs = Presentation(file_path)

            # Metadata
            props = prs.core_properties
            result["metadata"] = {
                "author": props.author,
                "created": str(props.created) if props.created else None,
                "modified": str(props.modified) if props.modified else None,
                "last_modified_by": props.last_modified_by,
                "title": props.title,
                "subject": props.subject,
                "keywords": props.keywords,
            }

            # Slides
            result["slides_count"] = len(prs.slides)
            text_parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
            result["text"] = "\n".join(text_parts)

        except Exception as e:
            result["error"] = str(e)

        return result